"""Caption-below image renderer with proportional fit to the text column."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image
from docx.shared import Emu, Length, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from src import scanner
from src.schema import ImageModel
from src.style_cloner import DEFAULT_CJK_FONT, set_run_fonts

_CAPTION_SIZE_PT = 10.5


class ImageNotFoundError(FileNotFoundError):
    """Raised when the payload points at a missing or unreadable image."""


@dataclass
class ImageRenderResult:
    """Handles to the inserted picture paragraph and its caption."""

    picture_paragraph: Paragraph
    caption: Paragraph
    width: Length
    last_element: Paragraph


def _resolved(path: str | Path) -> Path:
    resolved = Path(path).expanduser()
    if not resolved.is_file():
        raise ImageNotFoundError(f"图片不存在或不可读：{resolved.resolve()}")
    return resolved


def intrinsic_size(path: str | Path) -> tuple[Length, Length, int, int]:
    """Return ``(width_emu, height_emu, px_width, px_height)`` for an image."""

    resolved = _resolved(path)
    try:
        image = Image.from_file(str(resolved))
    except Exception as exc:  # unsupported/corrupt image data
        raise ImageNotFoundError(f"无法读取图片：{resolved}") from exc

    pixel_width = int(getattr(image, "px_width", 0) or 0)
    pixel_height = int(getattr(image, "px_height", 0) or 0)
    if pixel_width <= 0 or pixel_height <= 0:
        try:
            from PIL import Image as PilImage

            with PilImage.open(resolved) as pil_image:
                pixel_width, pixel_height = pil_image.size
        except Exception:
            pixel_width, pixel_height = 0, 0

    return Length(image.width), Length(image.height), pixel_width, pixel_height


def fit_image_size(path: str | Path, max_width: Length) -> tuple[Length, Length]:
    """Fit proportionally inside ``max_width``; never upscale."""

    width, height, _, _ = intrinsic_size(path)
    limit = int(max_width)
    if limit <= 0 or int(width) <= limit:
        return width, height
    scale = limit / int(width)
    return Emu(limit), Emu(max(1, int(int(height) * scale)))


def render_image(
    doc: DocumentObject,
    model: ImageModel,
    *,
    container: Paragraph | Table,
) -> ImageRenderResult:
    """Insert a centered picture, then its caption below it."""

    resolved = _resolved(model.path)
    column_width = scanner.content_width(doc)
    allowed = Emu(max(1, int(int(column_width) * model.max_width_ratio)))
    width, height = fit_image_size(resolved, allowed)

    picture_paragraph = scanner.insert_paragraph_after(doc, container)
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_format = picture_paragraph.paragraph_format
    picture_format.space_before = Pt(0)
    picture_format.space_after = Pt(0)
    picture_format.keep_with_next = True
    run = picture_paragraph.add_run()
    run.add_picture(str(resolved), width=width, height=height)

    caption = scanner.insert_paragraph_after(doc, picture_paragraph, model.title)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_format = caption.paragraph_format
    caption_format.space_before = Pt(3)
    caption_format.space_after = Pt(6)
    caption_format.line_spacing = 1.0
    caption_run = caption.runs[0] if caption.runs else caption.add_run()
    set_run_fonts(caption_run, DEFAULT_CJK_FONT)
    caption_run.font.size = Pt(_CAPTION_SIZE_PT)
    caption_run.font.bold = False

    return ImageRenderResult(
        picture_paragraph=picture_paragraph,
        caption=caption,
        width=width,
        last_element=caption,
    )


__all__ = [
    "ImageNotFoundError",
    "ImageRenderResult",
    "fit_image_size",
    "intrinsic_size",
    "render_image",
]
