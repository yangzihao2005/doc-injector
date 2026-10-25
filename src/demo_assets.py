"""Generate a realistic DSP lab-report template plus demo payload and image."""

from __future__ import annotations

import json
import math
import random
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image as PilImage
from PIL import ImageDraw

from src.schema import ReportInjectionPayload

_CJK = "宋体"
_LATIN = "Times New Roman"
_HEI = "黑体"


def _set_style_font(
    style: Any,
    *,
    latin: str,
    east_asia: str,
    size_pt: float,
    bold: bool | None = None,
) -> None:
    style.font.name = latin
    style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:cs"), latin)


def _set_run_font(run: Any, *, size_pt: float, bold: bool = False, east_asia: str = _CJK) -> None:
    run.font.name = _LATIN
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), _LATIN)
    r_fonts.set(qn("w:hAnsi"), _LATIN)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:cs"), _LATIN)


def _add_page_field(paragraph: Any) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)


def draw_waveform(path: str | Path) -> Path:
    """Draw a deterministic sine-plus-noise waveform PNG (no matplotlib)."""

    destination = Path(path).expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)

    width, height = 1200, 500
    image = PilImage.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)

    for x in range(0, width, 60):
        draw.line([(x, 0), (x, height)], fill=(226, 226, 226), width=1)
    for y in range(0, height, 50):
        draw.line([(0, y), (width, y)], fill=(226, 226, 226), width=1)

    mid = height // 2
    draw.line([(0, mid), (width, mid)], fill=(70, 70, 70), width=2)
    draw.line([(60, 0), (60, height)], fill=(70, 70, 70), width=2)

    random.seed(20240501)
    points: list[tuple[float, float]] = []
    for x in range(60, width - 20):
        phase = (x - 60) / 360.0
        value = 165.0 * math.sin(2.0 * math.pi * phase)
        value += random.uniform(-9.0, 9.0)
        points.append((float(x), mid - value))
    draw.line(points, fill=(0, 82, 155), width=2, joint="curve")

    image.save(destination, format="PNG")
    return destination


def build_sample_template(path: str | Path) -> Path:
    """Create a Chinese DSP lab-report template with anchors and rich styles."""

    destination = Path(path).expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)

    _set_style_font(doc.styles["Normal"], latin=_LATIN, east_asia=_CJK, size_pt=12, bold=False)
    normal_format = doc.styles["Normal"].paragraph_format
    normal_format.line_spacing = 1.5
    normal_format.space_before = Pt(0)
    normal_format.space_after = Pt(0)
    normal_format.first_line_indent = Cm(0.74)

    _set_style_font(
        doc.styles["Heading 1"], latin=_HEI, east_asia=_HEI, size_pt=14, bold=True
    )
    heading1 = doc.styles["Heading 1"].paragraph_format
    heading1.space_before = Pt(12)
    heading1.space_after = Pt(6)
    heading1.first_line_indent = Cm(0)
    heading1.line_spacing = 1.5

    _set_style_font(
        doc.styles["Heading 2"], latin=_HEI, east_asia=_HEI, size_pt=12, bold=True
    )
    heading2 = doc.styles["Heading 2"].paragraph_format
    heading2.space_before = Pt(6)
    heading2.space_after = Pt(3)
    heading2.first_line_indent = Cm(0)

    header_paragraph = section.header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header_paragraph.add_run("DSP 实验报告")
    _set_run_font(header_run, size_pt=9)

    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prefix = footer_paragraph.add_run("第 ")
    _set_run_font(prefix, size_pt=9)
    _add_page_field(footer_paragraph)
    suffix = footer_paragraph.add_run(" 页")
    _set_run_font(suffix, size_pt=9)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_format = title.paragraph_format
    title_format.space_before = Pt(6)
    title_format.space_after = Pt(6)
    title_format.first_line_indent = Cm(0)
    title_run = title.add_run("数字信号处理实验报告")
    _set_run_font(title_run, size_pt=18, bold=True, east_asia=_HEI)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle.paragraph_format.first_line_indent = Cm(0)
    subtitle_run = subtitle.add_run("实验二：离散时间信号的频域分析")
    _set_run_font(subtitle_run, size_pt=12, bold=False)

    doc.add_heading("一、实验目的", level=1)
    doc.add_paragraph(
        "掌握离散时间信号的频谱分析方法，理解 DFT 与 FFT 的数学关系，"
        "能够利用 DSP 实验平台完成信号采样、频谱计算与结果解释。"
    )
    doc.add_paragraph(
        "熟悉实验箱的信号源、示波器与上位机软件的基本操作，建立从时域波形到频域谱线的直观认识。"
    )

    doc.add_heading("二、实验原理", level=1)
    doc.add_paragraph(
        "对有限长离散序列进行离散傅里叶变换（DFT），可以得到其在单位圆上的等间隔频率采样。"
        "当序列长度为 2 的整数次幂时，采用按时间抽取的基-2 FFT 算法可将运算量由 O(N²) "
        "降至 O(N log N)。"
    )
    doc.add_paragraph(
        "频谱分辨率与采样频率、采样点数有关。为避免频谱混叠，采样频率必须高于信号最高频率的两倍。"
    )

    doc.add_heading("三、实验内容", level=1)
    doc.add_paragraph(
        "使用信号发生器产生不同频率的正弦信号，经 A/D 采样后送入 DSP 处理单元，"
        "分别记录采样序列与 FFT 幅频特性，并比较理论值与实测值。"
    )

    existing_table = doc.add_table(rows=2, cols=3)
    existing_table.style = "Table Grid"
    existing_headers = ("实验项目", "信号频率", "采样点数")
    existing_values = ("正弦信号采样", "1 kHz", "1024")
    for index, text in enumerate(existing_headers):
        cell_run = existing_table.cell(0, index).paragraphs[0].add_run(text)
        _set_run_font(cell_run, size_pt=10.5, bold=True)
    for index, text in enumerate(existing_values):
        cell_run = existing_table.cell(1, index).paragraphs[0].add_run(text)
        _set_run_font(cell_run, size_pt=10.5)

    doc.add_heading("四、实验结果分析", level=1)
    doc.add_paragraph(
        "将采集到的时域序列绘制成波形，并对 FFT 结果取模得到幅频特性曲线。"
        "观察主瓣宽度、旁瓣电平以及栅栏效应，分析采样参数对频谱的影响。"
    )

    doc.add_heading("五、实验总结", level=1)
    doc.add_paragraph(
        "通过本次实验加深了对采样定理与 FFT 算法的理解，掌握了频谱分析的基本流程，"
        "并认识到参数选择对测量结果的直接影响。"
    )

    doc.save(str(destination))
    return destination


def build_sample_payload(image_path: str | Path, path: str | Path) -> Path:
    """Write a demo payload that exercises paragraphs, both table styles and an image."""

    destination = Path(path).expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    resolved_image = Path(image_path).expanduser().resolve()

    raw: dict[str, Any] = {
        "blocks": [
            {
                "anchor_keyword": "实验目的",
                "paragraphs": [
                    "本次实验围绕离散时间信号的频域分析展开，目标是掌握 DFT/FFT 的基本使用方法。",
                    "实验过程中需要记录采样参数、观察频谱泄漏现象，并对误差来源做出解释。",
                ],
            },
            {
                "anchor_keyword": "实验原理",
                "paragraphs": [
                    "设采样频率为 fs，采样点数为 N，则频率分辨率为 fs/N。",
                ],
                "table": {
                    "title": "表2-1 典型离散时间信号",
                    "table_type": "three_line",
                    "headers": ["信号类型", "表达式", "主要特点"],
                    "rows": [
                        ["单位阶跃", "u(n)", "单边、含直流分量"],
                        ["指数序列", "a^n u(n)", "收敛性由 |a| 决定"],
                        ["正弦序列", "sin(ωn)", "周期性与ω有关"],
                    ],
                    "alignment": "center",
                },
            },
            {
                "anchor_keyword": "实验结果分析",
                "paragraphs": [
                    "下图给出了 1 kHz 正弦信号经采样后的时域波形，采样点数为 1024，"
                    "采样频率为 8 kHz。",
                ],
                "table": {
                    "title": "表4-1 采样测试数据",
                    "table_type": "grid",
                    "headers": ["序号", "信号频率/Hz", "采样值"],
                    "rows": [
                        [1, 500, 0.707],
                        [2, 1000, 0.991],
                        [3, 2000, 0.123],
                    ],
                    "alignment": "center",
                },
                "images": [
                    {
                        "title": "图4-1 正弦信号时域波形",
                        "path": str(resolved_image),
                        "max_width_ratio": 0.8,
                    }
                ],
            },
        ]
    }

    payload = ReportInjectionPayload.model_validate(raw)
    destination.write_text(
        json.dumps(payload.model_dump(mode="json"), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return destination


def build_demo_assets(directory: str | Path) -> dict[str, Path]:
    """Create template, image and payload inside ``directory``."""

    target = Path(directory).expanduser().resolve()
    target.mkdir(parents=True, exist_ok=True)
    template = build_sample_template(target / "sample_dsp.docx")
    image = draw_waveform(target / "waveform.png")
    payload = build_sample_payload(image, target / "sample_payload.json")
    return {"template": template, "image": image, "payload": payload}


__all__ = [
    "build_demo_assets",
    "build_sample_payload",
    "build_sample_template",
    "draw_waveform",
]
