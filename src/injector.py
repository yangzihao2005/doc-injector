"""Injection orchestrator: anchor lookup, ordered rendering, safe save, fallback.

The injector never rebuilds the document.  It opens the template, moves freshly
created paragraphs/tables/images next to the requested anchors, and writes to a
temporary file before atomically replacing the destination.  If the document or
the destination cannot be written (read-only, locked by Word, corrupt package),
the payload is exported as structured text and copied to the clipboard.
"""

from __future__ import annotations

import os
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table
from docx.text.paragraph import Paragraph

from src import scanner
from src.image_engine import ImageNotFoundError, render_image
from src.schema import ContentBlock, ReportInjectionPayload
from src.style_cloner import capture_body_style, write_styled_paragraph
from src.table_engine import render_table

Cursor = Paragraph | Table


class InjectionError(RuntimeError):
    """Base class for expected injection failures."""


class AnchorNotFoundError(InjectionError):
    """Raised when a payload block points at a keyword that is not present."""

    def __init__(self, keyword: str, available: list[str]) -> None:
        self.keyword = keyword
        self.available = available
        hint = "、".join(available[:8]) if available else "（模板中没有识别到大纲标题）"
        super().__init__(f"未找到锚点：{keyword!r}。可用大纲：{hint}")


@dataclass
class InjectionStats:
    """Counts and non-fatal warnings for one injection run."""

    blocks: int = 0
    paragraphs: int = 0
    tables: int = 0
    images: int = 0
    skipped_blocks: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass
class InjectionResult:
    """Outcome of an injection run, including fallback details when relevant."""

    template_path: Path
    output_path: Path
    stats: InjectionStats
    fallback_used: bool = False
    fallback_file: Path | None = None
    clipboard_text: str | None = None
    clipboard_copied: bool = False

    def summary(self) -> str:
        stats = self.stats
        line = (
            f"注入完成：{stats.blocks} 个内容块 / {stats.paragraphs} 段 / "
            f"{stats.tables} 表 / {stats.images} 图"
        )
        if self.fallback_used:
            line = f"注入降级：结构化文本已导出（{stats.blocks} 个内容块）"
        return f"{line}\n输出：{self.fallback_file or self.output_path}"


def _available_outline(doc: DocumentObject) -> list[str]:
    return [entry.text for entry in scanner.scan_outline(doc)]


def resolve_image_path(path: str, base_dir: Path | None = None) -> Path:
    """Resolve an image path, preferring ``base_dir`` for relative payload paths."""

    candidate = Path(path).expanduser()
    if not candidate.is_absolute() and base_dir is not None:
        candidate = base_dir / candidate
    return candidate.resolve()


def validate_payload_media(
    payload: ReportInjectionPayload, base_dir: Path | None = None
) -> dict[str, Path]:
    """Fail fast on missing media before the template is mutated."""

    resolved: dict[str, Path] = {}
    for block in payload.blocks:
        for image in block.images or []:
            path = resolve_image_path(image.path, base_dir)
            if not path.is_file():
                raise ImageNotFoundError(f"图片不存在或不可读：{path}")
            resolved[image.path] = path
    return resolved


def inject_blocks(
    doc: DocumentObject,
    blocks: list[ContentBlock],
    *,
    allow_missing_anchors: bool = False,
    base_dir: Path | None = None,
) -> InjectionStats:
    """Inject blocks into an already-open document (mutates ``doc``)."""

    stats = InjectionStats()
    anchor_cursors: dict[int, Cursor] = {}
    for block in blocks:
        anchor = scanner.find_anchor(doc, block.anchor_keyword)
        if anchor is None:
            if not allow_missing_anchors:
                raise AnchorNotFoundError(block.anchor_keyword, _available_outline(doc))
            message = f"跳过缺失锚点：{block.anchor_keyword!r}"
            stats.skipped_blocks.append(block.anchor_keyword)
            stats.warnings.append(message)
            continue

        snapshot = capture_body_style(doc, anchor)
        anchor_key = id(anchor._p)
        cursor: Cursor = anchor_cursors.get(anchor_key, anchor)

        for text in block.paragraphs or []:
            if not text.strip():
                continue
            paragraph = scanner.insert_paragraph_after(doc, cursor)
            write_styled_paragraph(
                paragraph,
                text,
                snapshot,
                font=block.custom_font,
                size_pt=block.custom_size_pt,
            )
            cursor = paragraph
            stats.paragraphs += 1

        if block.table is not None:
            rendered = render_table(doc, block.table, container=cursor, snapshot=snapshot)
            cursor = rendered.last_element
            stats.tables += 1

        for image in block.images or []:
            resolved = resolve_image_path(image.path, base_dir)
            if resolved == Path(image.path).resolve():
                image_for_render = image
            else:
                image_for_render = image.model_copy(update={"path": str(resolved)})
            rendered_image = render_image(doc, image_for_render, container=cursor)
            cursor = rendered_image.last_element
            stats.images += 1

        stats.blocks += 1
        anchor_cursors[anchor_key] = cursor
    return stats


def inject_report(
    payload: ReportInjectionPayload,
    template_path: str | Path,
    output_path: str | Path | None = None,
    *,
    allow_missing_anchors: bool = False,
    base_dir: str | Path | None = None,
    copy_on_fallback: bool = True,
) -> InjectionResult:
    """Inject ``payload`` into ``template_path`` and write ``output_path``.

    The template is only ever read.  A locking problem on the destination
    triggers the text/clipboard fallback instead of corrupting either file.
    """

    template = Path(template_path).expanduser().resolve()
    if not template.is_file():
        raise InjectionError(f"模板不存在：{template}")

    output = (
        Path(output_path).expanduser().resolve()
        if output_path is not None
        else template.with_name(f"{template.stem}-injected{template.suffix}")
    )

    media_base = Path(base_dir).resolve() if base_dir is not None else None
    validate_payload_media(payload, media_base)

    try:
        doc = Document(str(template))
    except Exception as exc:  # corrupt package / unsupported format
        return _fallback_result(
            payload,
            template,
            output,
            reason=f"无法打开模板：{exc}",
            stats=InjectionStats(),
            copy_on_fallback=copy_on_fallback,
        )

    stats = inject_blocks(
        doc,
        payload.blocks,
        allow_missing_anchors=allow_missing_anchors,
        base_dir=media_base,
    )

    try:
        _save_atomically(doc, output)
    except OSError as exc:
        return _fallback_result(
            payload,
            template,
            output,
            reason=f"保存失败（文件可能被占用或只读）：{exc}",
            stats=stats,
            copy_on_fallback=copy_on_fallback,
        )

    return InjectionResult(template_path=template, output_path=output, stats=stats)


def _save_atomically(doc: DocumentObject, output: Path) -> None:
    """Save to a sibling temp file, then atomically replace the destination."""

    output.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=f".{output.stem}-", suffix=".docx", dir=output.parent)
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        doc.save(str(temp_path))
        os.replace(temp_path, output)
    finally:
        temp_path.unlink(missing_ok=True)


def render_fallback_text(
    payload: ReportInjectionPayload,
    *,
    template_path: Path | None = None,
    reason: str | None = None,
) -> str:
    """Render the payload as clipboard-friendly structured text."""

    lines: list[str] = ["# doc-injector 结构化导出"]
    if template_path is not None:
        lines.append(f"模板：{template_path}")
    if reason:
        lines.append(f"降级原因：{reason}")
    lines.append("")

    for index, block in enumerate(payload.blocks, start=1):
        lines.append(f"## {index}. 锚点：{block.anchor_keyword}")
        lines.append("")
        for paragraph in block.paragraphs or []:
            if paragraph.strip():
                lines.append(paragraph.strip())
                lines.append("")

        if block.table is not None:
            table = block.table
            lines.append(table.title)
            lines.append("")
            lines.append("| " + " | ".join(table.headers) + " |")
            lines.append("| " + " | ".join("---" for _ in table.headers) + " |")
            for row in table.rows:
                lines.append("| " + " | ".join(str(cell) for cell in row) + " |")
            lines.append("")

        for image in block.images or []:
            lines.append(f"![{image.title}]({image.path})")
            lines.append(image.title)
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def _fallback_result(
    payload: ReportInjectionPayload,
    template: Path,
    output: Path,
    *,
    reason: str,
    stats: InjectionStats,
    copy_on_fallback: bool,
) -> InjectionResult:
    text = render_fallback_text(payload, template_path=template, reason=reason)
    output.parent.mkdir(parents=True, exist_ok=True)
    fallback_file = output.with_suffix(".fallback.txt")
    fallback_file.write_text(text, encoding="utf-8")
    copied = copy_to_clipboard(text) if copy_on_fallback else False
    stats.warnings.append(reason)
    return InjectionResult(
        template_path=template,
        output_path=output,
        stats=stats,
        fallback_used=True,
        fallback_file=fallback_file,
        clipboard_text=text,
        clipboard_copied=copied,
    )


def copy_to_clipboard(text: str) -> bool:
    """Best-effort clipboard copy via Tk; returns False on headless systems."""

    try:
        import tkinter

        root = tkinter.Tk()
        root.withdraw()
        root.clipboard_clear()
        root.clipboard_append(text)
        root.update()
        root.destroy()
        return True
    except Exception:
        return False


__all__ = [
    "AnchorNotFoundError",
    "Cursor",
    "InjectionError",
    "InjectionResult",
    "InjectionStats",
    "copy_to_clipboard",
    "inject_blocks",
    "inject_report",
    "render_fallback_text",
    "resolve_image_path",
    "validate_payload_media",
]
