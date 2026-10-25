"""Click-based command line interface for doc-injector."""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

import click
from pydantic import ValidationError

from src import scanner
from src.demo_assets import build_demo_assets
from src.image_engine import ImageNotFoundError
from src.injector import AnchorNotFoundError, InjectionError, inject_report
from src.schema import load_payload, payload_json_schema

try:  # pragma: no cover - exercised through the installed entry point
    from src import __version__
except ImportError:  # pragma: no cover
    __version__ = "0.0.0"


def _document_summary(template: Path) -> dict[str, Any]:
    from docx import Document

    doc = Document(str(template))
    width = scanner.content_width(doc)
    outline = [
        {
            "index": entry.index,
            "text": entry.text,
            "style": entry.style_name,
            "level": entry.level,
        }
        for entry in scanner.scan_outline(doc)
    ]
    anchors = [
        paragraph.text.strip()
        for paragraph in doc.paragraphs
        if paragraph.text.strip() and not scanner.is_heading(paragraph)
    ]
    return {
        "template": str(template),
        "sections": len(doc.sections),
        "content_width_cm": round(width.cm, 2),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "outline": outline,
        "anchor_candidates": anchors[:15],
    }


@click.group(context_settings={"help_option_names": ["-h", "--help"]})
@click.version_option(version=__version__, prog_name="doc-injector")
def main() -> None:
    """Surgically inject academic paragraphs, tables and figures into .docx."""


@main.command("inspect")
@click.argument("template", type=click.Path(exists=True, dir_okay=False, path_type=Path))
@click.option(
    "--json-output",
    "-j",
    type=click.Path(dir_okay=False, path_type=Path),
    help="Write the inspection result as JSON.",
)
def inspect_command(template: Path, json_output: Path | None) -> None:
    """List the template outline and candidate anchors."""

    try:
        summary = _document_summary(template.resolve())
    except Exception as exc:
        raise click.ClickException(f"无法读取模板：{exc}") from exc

    if json_output is not None:
        json_output.parent.mkdir(parents=True, exist_ok=True)
        json_output.write_text(
            json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        click.echo(f"已写入：{json_output}")
        return

    click.echo(f"模板：{summary['template']}")
    click.echo(
        f"节数：{summary['sections']}　版心宽度：{summary['content_width_cm']} cm　"
        f"段落：{summary['paragraphs']}　表格：{summary['tables']}"
    )
    click.echo("大纲：")
    outline = summary["outline"]
    if not outline:
        click.echo("  （未识别到标题样式，可改用正文关键词锚定）")
    for entry in outline:
        click.echo(f"  [{entry['index']:>3}] {entry['style']:<12} {entry['text']}")
    click.echo("可锚定正文（前 15 条）：")
    for text in summary["anchor_candidates"]:
        preview = text if len(text) <= 48 else text[:48] + "…"
        click.echo(f"  - {preview}")


@main.command("inject")
@click.argument("template", type=click.Path(exists=True, dir_okay=False, path_type=Path))
@click.argument("payload", type=click.Path(exists=True, dir_okay=False, path_type=Path))
@click.option(
    "--output",
    "-o",
    type=click.Path(dir_okay=False, path_type=Path),
    help="Output .docx path (default: <template>-injected.docx).",
)
@click.option(
    "--allow-missing",
    is_flag=True,
    help="Skip blocks whose anchor keyword is absent instead of failing.",
)
@click.option("--no-clipboard", is_flag=True, help="Do not touch the clipboard on fallback.")
def inject_command(
    template: Path,
    payload: Path,
    output: Path | None,
    allow_missing: bool,
    no_clipboard: bool,
) -> None:
    """Inject a JSON payload into a .docx template."""

    try:
        report = load_payload(payload)
    except (ValidationError, json.JSONDecodeError, OSError) as exc:
        raise click.ClickException(f"payload 校验失败：{exc}") from exc

    try:
        result = inject_report(
            report,
            template,
            output,
            allow_missing_anchors=allow_missing,
            base_dir=payload.resolve().parent,
            copy_on_fallback=not no_clipboard,
        )
    except AnchorNotFoundError as exc:
        raise click.ClickException(str(exc)) from exc
    except (InjectionError, ImageNotFoundError) as exc:
        raise click.ClickException(str(exc)) from exc

    for warning in result.stats.warnings:
        click.echo(f"警告：{warning}", err=True)
    click.echo(result.summary())
    if result.fallback_used and result.clipboard_copied:
        click.echo("结构化文本已复制到剪贴板。")


@main.command("demo")
@click.option(
    "--directory",
    "-d",
    type=click.Path(file_okay=False, path_type=Path),
    default=Path("examples"),
    show_default=True,
    help="Directory for the generated template, image, payload and output.",
)
@click.option("--no-inject", "skip_inject", is_flag=True, help="Only generate assets.")
def demo_command(directory: Path, skip_inject: bool) -> None:
    """Generate demo assets and, by default, run a full injection."""

    assets = build_demo_assets(directory)
    click.echo(f"模板：{assets['template']}")
    click.echo(f"波形图：{assets['image']}")
    click.echo(f"payload：{assets['payload']}")
    if skip_inject:
        return

    try:
        report = load_payload(assets["payload"])
        result = inject_report(
            report,
            assets["template"],
            assets["template"].with_name("sample_dsp-injected.docx"),
            base_dir=assets["payload"].parent,
        )
    except (InjectionError, ImageNotFoundError, ValidationError) as exc:
        raise click.ClickException(str(exc)) from exc

    click.echo(result.summary())


@main.command("schema")
@click.option(
    "--output",
    "-o",
    type=click.Path(dir_okay=False, path_type=Path),
    help="Write the JSON schema to a file instead of stdout.",
)
def schema_command(output: Path | None) -> None:
    """Print the payload JSON Schema (useful for editor autocomplete)."""

    text = json.dumps(payload_json_schema(), ensure_ascii=False, indent=2) + "\n"
    if output is None:
        click.echo(text)
        return
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(text, encoding="utf-8")
    click.echo(f"已写入：{output}")


if __name__ == "__main__":  # pragma: no cover
    sys.exit(main())
