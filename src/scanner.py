"""Template outline scanning and low-level OpenXML insertion helpers.

Anchors are located by normalized keyword containment first (headings win over
body text), then by a punctuation-tolerant fuzzy regex.  All insertion helpers
move newly created body elements next to an existing anchor instead of
appending them at the end of the document, which is what keeps untouched
content byte-for-byte stable.
"""

from __future__ import annotations

import re
import unicodedata
from collections.abc import Iterator
from dataclasses import dataclass

from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.shared import Emu, Length
from docx.table import Table
from docx.text.paragraph import Paragraph

Container = Paragraph | Table

_HEADING_STYLE_RE = re.compile(r"^(?:heading|标题)\s*(\d+)?$", re.IGNORECASE)
_CN_HEADING_RE = re.compile(r"^[一二三四五六七八九十]+[、.．]")
_CAPTION_STYLE_RE = re.compile(r"(caption|题注|图表标题)", re.IGNORECASE)
_PUNCTUATION_RE = re.compile(r"[\s\u3000:：、,，.．;；·\-—_()（）\[\]【】<>《》\"'“”‘’/\\]+")
_DRAWING_XPATH = f".//{qn('w:drawing')}"
_DEFAULT_WIDTH = Emu(int(6.0 * 914400))


@dataclass(frozen=True)
class OutlineEntry:
    """A body paragraph that looks like an outline / heading."""

    index: int
    text: str
    style_name: str
    level: int | None


def normalize_text(text: str) -> str:
    """NFKC-normalize, strip punctuation/whitespace and casefold for matching."""

    normalized = unicodedata.normalize("NFKC", text)
    return _PUNCTUATION_RE.sub("", normalized).casefold()


def flexible_pattern(keyword: str) -> re.Pattern[str] | None:
    """Build a regex tolerant of spaces/punctuation between keyword characters."""

    chars = [char for char in unicodedata.normalize("NFKC", keyword) if not char.isspace()]
    if not chars:
        return None
    body = r"[\s\W_]*".join(re.escape(char) for char in chars)
    return re.compile(body, re.IGNORECASE)


def heading_level(paragraph: Paragraph) -> int | None:
    """Return a 1-based heading level for a paragraph, or ``None``."""

    style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
    match = _HEADING_STYLE_RE.match(style_name.strip())
    if match is not None:
        return int(match.group(1) or 1)

    p_pr = paragraph._p.pPr
    if p_pr is not None:
        outline = p_pr.find(qn("w:outlineLvl"))
        if outline is not None:
            raw = outline.get(qn("w:val"))
            if raw is not None and raw.isdigit():
                return int(raw) + 1

    # Many lab-report templates leave section headings as "Normal" and rely on
    # a Chinese numeral prefix instead of a Heading style.
    if _CN_HEADING_RE.match(paragraph.text.strip()):
        return 1
    return None


def is_caption(paragraph: Paragraph) -> bool:
    style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
    return _CAPTION_STYLE_RE.search(style_name) is not None


def is_heading(paragraph: Paragraph) -> bool:
    return heading_level(paragraph) is not None


def iter_paragraphs(doc: DocumentObject) -> Iterator[Paragraph]:
    """Iterate every paragraph in document order, including nested table cells.

    Top-level body paragraphs keep their natural order.  When a table is
    encountered, its rows/cells are walked depth-first (including tables nested
    inside cells) so anchors like "一、实验目的" inside a report shell table are
    visible to the rest of the scanner.
    """

    yield from _iter_block_paragraphs(doc.element.body, doc._body)


def _iter_block_paragraphs(container: BaseOxmlElement, parent: object) -> Iterator[Paragraph]:
    for child in container.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)  # type: ignore[arg-type]
        elif child.tag == qn("w:tbl"):
            yield from _iter_table_paragraphs(child, parent)
        elif child.tag == qn("w:sdt"):
            content = child.find(qn("w:sdtContent"))
            if content is not None:
                yield from _iter_block_paragraphs(content, parent)


def _iter_table_paragraphs(table_element: BaseOxmlElement, parent: object) -> Iterator[Paragraph]:
    for row in table_element.iterchildren(qn("w:tr")):
        for cell in row.iterchildren(qn("w:tc")):
            yield from _iter_block_paragraphs(cell, parent)


def scan_outline(doc: DocumentObject) -> list[OutlineEntry]:
    """Return every heading-like paragraph, including Chinese-numbered ones."""

    entries: list[OutlineEntry] = []
    for index, paragraph in enumerate(iter_paragraphs(doc)):
        text = paragraph.text.strip()
        if not text:
            continue
        level = heading_level(paragraph)
        if level is None:
            continue
        style_name = paragraph.style.name if paragraph.style is not None else ""
        entries.append(OutlineEntry(index=index, text=text, style_name=style_name, level=level))
    return entries


def find_anchor(doc: DocumentObject, keyword: str) -> Paragraph | None:
    """Find the best paragraph for ``keyword``.

    Priority: heading containing the normalized keyword, body paragraph
    containing it, then fuzzy regex over raw text.  First match wins, which
    mirrors the top-down reading order of a report.
    """

    target = normalize_text(keyword)
    if not target:
        return None

    paragraphs = list(iter_paragraphs(doc))

    for paragraph in paragraphs:
        if is_heading(paragraph) and target in normalize_text(paragraph.text):
            return paragraph

    for paragraph in paragraphs:
        if target in normalize_text(paragraph.text):
            return paragraph

    pattern = flexible_pattern(keyword)
    if pattern is not None:
        for paragraph in paragraphs:
            if pattern.search(paragraph.text):
                return paragraph
    return None


def find_section_end(doc: DocumentObject, anchor: Paragraph) -> Paragraph | None:
    """Return the next heading at the same-or-higher level after ``anchor``."""

    start_level = heading_level(anchor)
    seen = False
    for paragraph in iter_paragraphs(doc):
        if paragraph._p is anchor._p:
            seen = True
            continue
        if not seen:
            continue
        level = heading_level(paragraph)
        if level is None:
            continue
        if start_level is None or level <= start_level:
            return paragraph
    return None


def find_body_sample(doc: DocumentObject, anchor: Paragraph) -> Paragraph | None:
    """Return the first plain body paragraph after ``anchor`` within its section."""

    section_end = find_section_end(doc, anchor)
    seen = False
    for paragraph in iter_paragraphs(doc):
        if paragraph._p is anchor._p:
            seen = True
            continue
        if not seen:
            continue
        if section_end is not None and paragraph._p is section_end._p:
            break
        text = paragraph.text.strip()
        if not text or is_heading(paragraph) or is_caption(paragraph):
            continue
        if paragraph._p.find(_DRAWING_XPATH) is not None:
            continue
        return paragraph
    return None


def first_body_paragraph(doc: DocumentObject) -> Paragraph | None:
    """Fallback style source: first non-empty, non-heading body paragraph."""

    for paragraph in iter_paragraphs(doc):
        if not paragraph.text.strip() or is_heading(paragraph):
            continue
        return paragraph
    return None


def content_width(doc: DocumentObject) -> Length:
    """Effective text-column width of the first section: page width - margins."""

    sections = doc.sections
    if not sections:
        return _DEFAULT_WIDTH
    section = sections[0]
    if section.page_width is None or section.left_margin is None or section.right_margin is None:
        return _DEFAULT_WIDTH
    width = int(section.page_width) - int(section.left_margin) - int(section.right_margin)
    if width <= 0:
        return _DEFAULT_WIDTH
    return Emu(width)


def element_of(container: Container) -> BaseOxmlElement:
    """Return the underlying ``CT_P`` / ``CT_Tbl`` element of a wrapper."""

    if isinstance(container, Paragraph):
        return container._p
    return container._tbl


def insert_paragraph_after(
    doc: DocumentObject,
    container: Container,
    text: str = "",
    style: str | None = None,
) -> Paragraph:
    """Create a body paragraph immediately after ``container``."""

    paragraph = doc.add_paragraph(text, style=style)
    element_of(container).addnext(paragraph._p)
    return paragraph


def insert_table_after(
    doc: DocumentObject,
    container: Container,
    rows: int,
    cols: int,
) -> Table:
    """Create a table immediately after ``container``.

    When the anchor lives inside a table cell, Word expects a paragraph after a
    nested table.  A blank paragraph is appended only if the cell would
    otherwise end with the new table (or continue with another table).
    """

    container_element = element_of(container)
    table = doc.add_table(rows=rows, cols=cols)
    container_element.addnext(table._tbl)

    parent = container_element.getparent()
    if parent is not None and parent.tag == qn("w:tc"):
        following = table._tbl.getnext()
        if following is None or following.tag != qn("w:p"):
            filler = doc.add_paragraph()
            table._tbl.addnext(filler._p)
    return table


def section_signature(doc: DocumentObject) -> dict[str, object]:
    """Stable snapshot of layout properties used by the preservation tests."""

    signature: dict[str, object] = {}
    for index, section in enumerate(doc.sections):
        prefix = f"section{index}"
        signature[f"{prefix}.page_width"] = int(section.page_width or 0)
        signature[f"{prefix}.page_height"] = int(section.page_height or 0)
        signature[f"{prefix}.left_margin"] = int(section.left_margin or 0)
        signature[f"{prefix}.right_margin"] = int(section.right_margin or 0)
        signature[f"{prefix}.top_margin"] = int(section.top_margin or 0)
        signature[f"{prefix}.bottom_margin"] = int(section.bottom_margin or 0)
        signature[f"{prefix}.header_distance"] = int(section.header_distance or 0)
        signature[f"{prefix}.footer_distance"] = int(section.footer_distance or 0)
        signature[f"{prefix}.header_text"] = "\n".join(
            paragraph.text for paragraph in section.header.paragraphs
        )
        signature[f"{prefix}.footer_text"] = "\n".join(
            paragraph.text for paragraph in section.footer.paragraphs
        )
    return signature
