"""Shared fixtures: styled template, waveform image and a realistic payload."""

from __future__ import annotations

import math
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image as PilImage
from PIL import ImageDraw

from src.schema import ReportInjectionPayload


def _set_run_style(run: object, *, font: str, size_pt: float, bold: bool = False) -> None:
    run.font.name = font  # type: ignore[attr-defined]
    run.font.size = Pt(size_pt)  # type: ignore[attr-defined]
    run.font.bold = bold  # type: ignore[attr-defined]
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)  # type: ignore[attr-defined]
    r_pr = run._r.get_or_add_rPr()  # type: ignore[attr-defined]
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), font)
    r_fonts.set(qn("w:hAnsi"), font)
    r_fonts.set(qn("w:eastAsia"), font)
    r_fonts.set(qn("w:cs"), font)


@pytest.fixture
def template_path(tmp_path: Path) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)

    header = section.header.paragraphs[0]
    header.text = "DSP 实验报告"
    footer = section.footer.paragraphs[0]
    footer.text = "第 1 页"

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5

    doc.add_paragraph("数字信号处理实验报告")
    doc.add_heading("一、实验目的", level=1)
    purpose = doc.add_paragraph()
    purpose.paragraph_format.line_spacing = 1.5
    _set_run_style(
        purpose.add_run("实验目的正文，用于样式克隆。"),
        font="楷体",
        size_pt=14,
        bold=True,
    )

    doc.add_heading("二、实验原理", level=1)
    doc.add_paragraph("实验原理正文，保持原样。")

    doc.add_heading("三、实验结果分析", level=1)
    doc.add_paragraph("结果分析正文，保持原样。")

    path = tmp_path / "template.docx"
    doc.save(str(path))
    return path


@pytest.fixture
def waveform_path(tmp_path: Path) -> Path:
    path = tmp_path / "wave.png"
    image = PilImage.new("RGB", (1600, 600), "white")
    draw = ImageDraw.Draw(image)
    points = [
        (x, 300 - 180 * math.sin(2 * math.pi * x / 400)) for x in range(0, 1600, 5)
    ]
    draw.line(points, fill=(0, 80, 160), width=3)
    image.save(path, format="PNG")
    return path


@pytest.fixture
def payload(waveform_path: Path) -> ReportInjectionPayload:
    return ReportInjectionPayload.model_validate(
        {
            "blocks": [
                {
                    "anchor_keyword": "实验目的",
                    "paragraphs": ["注入的第一段实验目的。", "注入的第二段实验目的。"],
                },
                {
                    "anchor_keyword": "实验原理",
                    "table": {
                        "title": "表2-1 典型离散时间信号",
                        "table_type": "three_line",
                        "headers": ["信号类型", "表达式", "主要特点"],
                        "rows": [
                            ["单位阶跃", "u(n)", "单边"],
                            ["正弦序列", "sin(ωn)", "周期性"],
                        ],
                        "alignment": "center",
                    },
                },
                {
                    "anchor_keyword": "实验结果分析",
                    "paragraphs": ["注入的结果分析段落。"],
                    "table": {
                        "title": "表4-1 采样测试数据",
                        "table_type": "grid",
                        "headers": ["序号", "采样值"],
                        "rows": [[1, 0.12], [2, 0.98]],
                    },
                    "images": [
                        {
                            "title": "图4-1 正弦信号时域波形",
                            "path": str(waveform_path),
                            "max_width_ratio": 0.8,
                        }
                    ],
                },
            ]
        }
    )
