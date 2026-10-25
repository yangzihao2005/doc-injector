"""End-to-end injection tests: preservation, ordering, styles and fallbacks."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from pydantic import ValidationError

import src.injector as injector_module
from src import scanner
from src.image_engine import ImageNotFoundError
from src.injector import AnchorNotFoundError, inject_report
from src.schema import ContentBlock, ImageModel, ReportInjectionPayload, TableModel


def _reopen(path: Path) -> DocumentObject:
    return Document(str(path))


def _body_children(doc: DocumentObject) -> list[Any]:
    return list(doc.element.body.iterchildren())


def _previous_text(doc: DocumentObject, element: Any) -> str:
    children = _body_children(doc)
    index = children.index(element)
    assert index > 0
    return Paragraph(children[index - 1], doc._body).text


def _paragraph_index(doc: DocumentObject, target: object) -> int:
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph._p is target:
            return index
    raise AssertionError("paragraph not found in document body")


def test_injection_preserves_layout_headers_and_untouched_paragraphs(
    payload: ReportInjectionPayload, template_path: Path, tmp_path: Path
) -> None:
    before = _reopen(template_path)
    before_signature = scanner.section_signature(before)
    before_texts = [paragraph.text for paragraph in before.paragraphs]

    output = tmp_path / "out.docx"
    result = inject_report(payload, template_path, output)

    assert not result.fallback_used
    assert output.is_file()
    assert result.stats.paragraphs == 3
    assert result.stats.tables == 2
    assert result.stats.images == 1
    assert result.stats.blocks == 3
    assert result.stats.warnings == []

    after = _reopen(output)
    assert scanner.section_signature(after) == before_signature
    after_texts = [paragraph.text for paragraph in after.paragraphs]
    assert [text for text in after_texts if text in before_texts] == before_texts
    assert "注入的第一段实验目的。" in after_texts
    assert "注入的结果分析段落。" in after_texts
    assert len(after.tables) == 2


def test_table_caption_above_and_image_caption_below(
    payload: ReportInjectionPayload, template_path: Path, tmp_path: Path
) -> None:
    output = tmp_path / "order.docx"
    inject_report(payload, template_path, output)
    doc = _reopen(output)

    grid_table = doc.tables[-1]
    assert _previous_text(doc, grid_table._tbl) == "表4-1 采样测试数据"

    image_paragraphs = [
        paragraph
        for paragraph in doc.paragraphs
        if paragraph._p.find(f".//{qn('w:drawing')}") is not None
    ]
    assert len(image_paragraphs) == 1
    image_paragraph = image_paragraphs[0]
    assert image_paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER

    image_index = _paragraph_index(doc, image_paragraph._p)
    caption = doc.paragraphs[image_index + 1]
    assert caption.text == "图4-1 正弦信号时域波形"
    assert caption.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert caption.paragraph_format.space_before == Pt(3)
    assert caption.paragraph_format.space_after == Pt(6)

    assert len(doc.inline_shapes) == 1
    allowed = int(scanner.content_width(doc) * 0.8)
    assert int(doc.inline_shapes[0].width) <= allowed + 1


def test_style_is_cloned_from_body_sample_and_can_be_overridden(
    template_path: Path, tmp_path: Path
) -> None:
    cloned = ReportInjectionPayload.model_validate(
        {
            "blocks": [
                {
                    "anchor_keyword": "实验目的",
                    "paragraphs": ["克隆样式段落"],
                }
            ]
        }
    )
    cloned_output = tmp_path / "cloned.docx"
    inject_report(cloned, template_path, cloned_output)
    cloned_doc = _reopen(cloned_output)
    cloned_paragraph = next(p for p in cloned_doc.paragraphs if p.text == "克隆样式段落")
    cloned_run = cloned_paragraph.runs[0]

    assert cloned_run.font.name == "楷体"
    assert cloned_run.font.size == Pt(14)
    assert cloned_run.font.bold is True
    assert cloned_run._r.rPr is not None
    assert cloned_run._r.rPr.rFonts is not None
    assert cloned_run._r.rPr.rFonts.get(qn("w:eastAsia")) == "楷体"
    assert cloned_paragraph.paragraph_format.line_spacing == 1.5

    overridden = ReportInjectionPayload.model_validate(
        {
            "blocks": [
                {
                    "anchor_keyword": "实验目的",
                    "paragraphs": ["覆盖样式段落"],
                    "custom_font": "黑体",
                    "custom_size_pt": 10.5,
                }
            ]
        }
    )
    overridden_output = tmp_path / "overridden.docx"
    inject_report(overridden, template_path, overridden_output)
    overridden_doc = _reopen(overridden_output)
    overridden_paragraph = next(
        p for p in overridden_doc.paragraphs if p.text == "覆盖样式段落"
    )
    overridden_run = overridden_paragraph.runs[0]

    assert overridden_run.font.name == "黑体"
    assert overridden_run.font.size == Pt(10.5)
    assert overridden_run._r.rPr is not None
    assert overridden_run._r.rPr.rFonts is not None
    assert overridden_run._r.rPr.rFonts.get(qn("w:eastAsia")) == "黑体"


def test_missing_anchor_raises_or_is_skipped(
    template_path: Path, tmp_path: Path
) -> None:
    missing = ReportInjectionPayload.model_validate(
        {
            "blocks": [
                {"anchor_keyword": "不存在的锚点", "paragraphs": ["内容"]},
            ]
        }
    )
    output = tmp_path / "missing.docx"
    with pytest.raises(AnchorNotFoundError) as error:
        inject_report(missing, template_path, output)
    assert "不存在的锚点" in str(error.value)
    assert not output.exists()

    result = inject_report(
        missing,
        template_path,
        output,
        allow_missing_anchors=True,
    )
    assert output.is_file()
    assert result.stats.skipped_blocks == ["不存在的锚点"]
    assert result.stats.blocks == 0
    assert any("不存在的锚点" in warning for warning in result.stats.warnings)


def test_locked_output_falls_back_to_structured_text(
    payload: ReportInjectionPayload,
    template_path: Path,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def _locked(doc: object, output: Path) -> None:
        raise PermissionError("locked by Word")

    monkeypatch.setattr(injector_module, "_save_atomically", _locked)
    output = tmp_path / "locked.docx"
    result = inject_report(payload, template_path, output, copy_on_fallback=False)

    assert result.fallback_used
    assert not output.exists()
    assert result.fallback_file is not None
    assert result.fallback_file.is_file()
    assert result.fallback_file == output.with_suffix(".fallback.txt")
    assert result.clipboard_copied is False
    assert result.clipboard_text is not None
    assert "表2-1 典型离散时间信号" in result.clipboard_text
    assert "图4-1 正弦信号时域波形" in result.clipboard_text
    assert result.stats.tables == 2


def test_corrupt_template_falls_back(
    payload: ReportInjectionPayload, tmp_path: Path
) -> None:
    broken = tmp_path / "broken.docx"
    broken.write_text("this is not a docx package", encoding="utf-8")
    output = tmp_path / "broken-out.docx"

    result = inject_report(payload, broken, output, copy_on_fallback=False)

    assert result.fallback_used
    assert result.fallback_file is not None
    assert "降级原因" in result.fallback_file.read_text(encoding="utf-8")
    assert not output.exists()


def test_missing_image_fails_before_writing(
    template_path: Path, tmp_path: Path
) -> None:
    payload = ReportInjectionPayload.model_validate(
        {
            "blocks": [
                {
                    "anchor_keyword": "实验目的",
                    "images": [{"title": "图1-1 缺失图片", "path": "nope.png"}],
                }
            ]
        }
    )
    template_before = template_path.read_bytes()
    output = tmp_path / "no-image.docx"

    with pytest.raises(ImageNotFoundError):
        inject_report(payload, template_path, output, base_dir=tmp_path)

    assert not output.exists()
    assert template_path.read_bytes() == template_before


def test_relative_image_path_is_resolved_against_base_dir(
    template_path: Path, waveform_path: Path, tmp_path: Path
) -> None:
    payload = ReportInjectionPayload.model_validate(
        {
            "blocks": [
                {
                    "anchor_keyword": "实验目的",
                    "images": [
                        {
                            "title": "图1-2 相对路径波形",
                            "path": waveform_path.name,
                        }
                    ],
                }
            ]
        }
    )
    output = tmp_path / "relative.docx"
    result = inject_report(payload, template_path, output, base_dir=tmp_path)

    assert not result.fallback_used
    doc = _reopen(output)
    assert len(doc.inline_shapes) == 1
    assert output.is_file()


def test_repeated_anchor_keeps_block_order(
    template_path: Path, tmp_path: Path
) -> None:
    payload = ReportInjectionPayload.model_validate(
        {
            "blocks": [
                {"anchor_keyword": "实验目的", "paragraphs": ["同名锚点第一块"]},
                {"anchor_keyword": "实验目的", "paragraphs": ["同名锚点第二块"]},
            ]
        }
    )
    output = tmp_path / "repeated.docx"
    inject_report(payload, template_path, output)

    doc = _reopen(output)
    texts = [paragraph.text for paragraph in doc.paragraphs]
    assert texts.index("同名锚点第一块") < texts.index("同名锚点第二块")


def test_payload_contract_rejects_bad_shapes() -> None:
    with pytest.raises(ValidationError):
        TableModel(title="表1", headers=["A", "B"], rows=[["only-one"]])

    with pytest.raises(ValidationError):
        ContentBlock(anchor_keyword="实验目的")

    with pytest.raises(ValidationError):
        ImageModel(title="图1", path="x.png", max_width_ratio=1.5)
