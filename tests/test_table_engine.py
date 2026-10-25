"""XML-level assertions for the three-line and full-grid table renderer."""

from __future__ import annotations

from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from src.schema import TableModel, TableType
from src.table_engine import render_table


def _three_line_model() -> TableModel:
    return TableModel(
        title="表1-1 采样测试数据",
        table_type=TableType.THREE_LINE,
        headers=["序号", "采样值", "备注"],
        rows=[[1, 0.12, "正常"], [2, 0.98, "正常"]],
        alignment="center",
    )


def _edge(cell: Any, name: str) -> Any:
    tc_pr = cell._tc.tcPr
    borders = tc_pr.find(qn("w:tcBorders"))
    assert borders is not None
    return borders.find(qn(f"w:{name}"))


def _edge_value(cell: Any, name: str, attribute: str) -> str | None:
    edge = _edge(cell, name)
    assert edge is not None
    return edge.get(qn(f"w:{attribute}"))


def test_caption_is_above_table_with_academic_spacing() -> None:
    doc = Document()
    anchor = doc.add_paragraph("anchor")
    result = render_table(doc, _three_line_model(), container=anchor)

    children = list(doc.element.body.iterchildren())
    assert children.index(result.caption._p) + 1 == children.index(result.table._tbl)
    assert result.caption.text == "表1-1 采样测试数据"
    assert result.caption.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert result.caption.paragraph_format.space_before == Pt(6)
    assert result.caption.paragraph_format.space_after == Pt(3)
    assert result.caption.paragraph_format.keep_with_next is True
    assert result.last_element is result.table

    caption_run = result.caption.runs[0]
    assert caption_run.font.size == Pt(10.5)
    assert caption_run._r.rPr is not None
    assert caption_run._r.rPr.rFonts is not None
    assert caption_run._r.rPr.rFonts.get(qn("w:eastAsia")) == "宋体"


def test_three_line_borders_are_correct() -> None:
    doc = Document()
    anchor = doc.add_paragraph("anchor")
    result = render_table(doc, _three_line_model(), container=anchor)
    table = result.table

    assert _edge_value(table.cell(0, 0), "top", "val") == "single"
    assert _edge_value(table.cell(0, 0), "top", "sz") == "12"
    assert _edge_value(table.cell(0, 0), "bottom", "val") == "single"
    assert _edge_value(table.cell(0, 0), "bottom", "sz") == "6"
    assert _edge_value(table.cell(0, 0), "left", "val") == "nil"
    assert _edge_value(table.cell(0, 0), "right", "val") == "nil"

    assert _edge_value(table.cell(1, 1), "top", "val") == "nil"
    assert _edge_value(table.cell(1, 1), "left", "val") == "nil"

    last_row = len(table.rows) - 1
    assert _edge_value(table.cell(last_row, 0), "bottom", "val") == "single"
    assert _edge_value(table.cell(last_row, 0), "bottom", "sz") == "12"

    tbl_borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    assert tbl_borders is not None
    inside_v = tbl_borders.find(qn("w:insideV"))
    assert inside_v is not None
    assert inside_v.get(qn("w:val")) == "nil"


def test_three_line_header_only_table_keeps_thick_bottom() -> None:
    doc = Document()
    anchor = doc.add_paragraph("anchor")
    model = TableModel(
        title="表1-2 单行表",
        table_type=TableType.THREE_LINE,
        headers=["唯一栏目"],
        rows=[],
    )
    result = render_table(doc, model, container=anchor)
    assert len(result.table.rows) == 1
    assert _edge_value(result.table.cell(0, 0), "top", "sz") == "12"
    assert _edge_value(result.table.cell(0, 0), "bottom", "sz") == "12"


def test_grid_borders_cover_every_cell_edge() -> None:
    doc = Document()
    anchor = doc.add_paragraph("anchor")
    model = TableModel(
        title="表3-1 全框网格表",
        table_type=TableType.GRID,
        headers=["A", "B"],
        rows=[[1, 2], [3, 4]],
    )
    result = render_table(doc, model, container=anchor)

    for row in result.table.rows:
        for cell in row.cells:
            for edge in ("top", "left", "bottom", "right"):
                assert _edge_value(cell, edge, "val") == "single"
                assert _edge_value(cell, edge, "sz") == "4"

    tbl_borders = result.table._tbl.tblPr.find(qn("w:tblBorders"))
    assert tbl_borders is not None
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = tbl_borders.find(qn(f"w:{edge}"))
        assert element is not None
        assert element.get(qn("w:val")) == "single"


def test_rows_do_not_split_and_header_repeats() -> None:
    doc = Document()
    anchor = doc.add_paragraph("anchor")
    result = render_table(doc, _three_line_model(), container=anchor)

    for index, row in enumerate(result.table.rows):
        tr_pr = row._tr.trPr
        assert tr_pr is not None
        assert tr_pr.find(qn("w:cantSplit")) is not None
        header = tr_pr.find(qn("w:tblHeader"))
        if index == 0:
            assert header is not None
        else:
            assert header is None


def test_cell_alignment_and_header_bold() -> None:
    doc = Document()
    anchor = doc.add_paragraph("anchor")
    model = TableModel(
        title="表1-3 左对齐表",
        headers=["栏目", "值"],
        rows=[["a", 1]],
        alignment="left",
    )
    result = render_table(doc, model, container=anchor)

    for row in result.table.rows:
        for cell in row.cells:
            assert cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT

    header_run = result.table.cell(0, 0).paragraphs[0].runs[0]
    body_run = result.table.cell(1, 0).paragraphs[0].runs[0]
    assert header_run.font.bold is True
    assert body_run.font.bold is False
