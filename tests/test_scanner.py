"""Scanner tests for nested-table outlines and cell-safe insertion."""

from __future__ import annotations

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.shared import Pt

from src.scanner import (
    find_anchor,
    insert_paragraph_after,
    insert_table_after,
    iter_paragraphs,
    scan_outline,
)
from src.style_cloner import capture_body_style


def _nested_table_doc() -> DocumentObject:
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    cell = outer.cell(0, 0)
    cell.paragraphs[0].text = "封面占位"
    cell.add_paragraph("一、实验目的")
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).paragraphs[0].text = "四、实验结果与分析"
    return doc


def test_scan_nested_table_paragraphs() -> None:
    doc = _nested_table_doc()

    texts = [paragraph.text for paragraph in iter_paragraphs(doc)]
    assert "一、实验目的" in texts
    assert "四、实验结果与分析" in texts

    purpose = find_anchor(doc, "实验目的")
    assert purpose is not None
    assert purpose.text.strip() == "一、实验目的"
    assert purpose._p.getparent() is not None
    assert purpose._p.getparent().tag == qn("w:tc")

    result = find_anchor(doc, "结果与分析")
    assert result is not None
    assert result.text.strip() == "四、实验结果与分析"


def test_inspect_outline_detects_cn_numbers() -> None:
    doc = Document()
    doc.add_paragraph("四、实验结果与分析")
    doc.add_heading("二、实验原理", level=2)

    entries = scan_outline(doc)
    by_text = {entry.text: entry for entry in entries}

    assert by_text["四、实验结果与分析"].level == 1
    assert by_text["四、实验结果与分析"].style_name == "Normal"
    assert by_text["二、实验原理"].level == 2


def test_style_fallback_skips_cover_title_when_section_has_no_body() -> None:
    doc = Document()
    cover = doc.add_paragraph()
    cover_run = cover.add_run("封面标题")
    cover_run.font.size = Pt(42)

    outer = doc.add_table(rows=1, cols=1)
    cell = outer.cell(0, 0)
    cell.paragraphs[0].text = "一、实验目的"
    cell.add_paragraph("二、实验内容")

    body = doc.add_paragraph()
    body_run = body.add_run("正文样式")
    body_run.font.name = "宋体"
    body_run.font.size = Pt(12)

    anchor = find_anchor(doc, "实验目的")
    assert anchor is not None
    snapshot = capture_body_style(doc, anchor)
    assert snapshot.run.size_pt == 12


def test_insert_after_cell_anchor_stays_in_cell_and_ends_with_paragraph() -> None:
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    anchor = outer.cell(0, 0).paragraphs[0]
    anchor.text = "一、实验目的"

    paragraph = insert_paragraph_after(doc, anchor, "注入段落")
    assert paragraph._p.getparent() is anchor._p.getparent()
    assert paragraph._p.getprevious() is anchor._p

    table = insert_table_after(doc, paragraph, rows=1, cols=1)
    parent = table._tbl.getparent()
    assert parent is not None
    assert parent.tag == qn("w:tc")

    following = table._tbl.getnext()
    assert following is not None
    assert following.tag == qn("w:p")
